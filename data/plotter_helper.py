from openpyxl import Workbook
from openpyxl.chart import (
    Reference,
    Series,
    BarChart3D,
    LineChart,
    BarChart
)

def create_plot(chart_type, raw_data, plot_title, x_axis_title, y_axis_title, output_dir):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()

    for row in raw_data:
        ws.append(row)

    data = Reference(ws, min_col=2, min_row=1, max_col=len(raw_data[0]), max_row=len(raw_data))
    titles = Reference(ws, min_col=1, min_row=2, max_row=len(raw_data))

    if chart_type == "bar3d":
        chart = BarChart3D()
    elif chart_type == "bar2d":
        chart = BarChart()
        chart.type = "col"
    else:
        chart = LineChart()

    chart.title = plot_title
    chart.x_axis.title = x_axis_title
    chart.y_axis.title = y_axis_title

    chart.add_data(data, titles_from_data=True)
    
    chart.set_categories(titles)

    ws.add_chart(chart, "F5")
    wb.save(output_dir)